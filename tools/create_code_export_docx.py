from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "artifacts" / "ZBX_Website_Code_2000_Characters.docx"
CHARACTER_LIMIT = 2000

SOURCE_FILES = [
    ROOT / "app" / "layout.tsx",
    ROOT / "app" / "page.tsx",
    ROOT / "styles" / "tokens.css",
    ROOT / "app" / "globals.css",
]

COLORS = {
    "ink": RGBColor(26, 26, 26),
    "muted": RGBColor(112, 116, 120),
    "blue": RGBColor(0, 92, 163),
    "green": RGBColor(34, 134, 58),
    "purple": RGBColor(111, 66, 193),
    "red": RGBColor(176, 50, 55),
    "orange": RGBColor(167, 92, 0),
    "comment": RGBColor(106, 115, 125),
}

KEYWORDS = {
    "as", "async", "await", "break", "case", "catch", "class", "const",
    "continue", "default", "else", "export", "extends", "false", "finally",
    "for", "from", "function", "if", "import", "in", "interface", "keyof",
    "let", "new", "null", "of", "return", "satisfies", "switch", "throw",
    "true", "try", "type", "typeof", "undefined", "while",
}

TOKEN_RE = re.compile(
    r"(?P<comment>//.*$|/\*.*?\*/)"
    r"|(?P<string>`(?:\\.|[^`])*`|'(?:\\.|[^'])*'|\"(?:\\.|[^\"])*\")"
    r"|(?P<number>\b\d+(?:\.\d+)?\b)"
    r"|(?P<word>\b[A-Za-z_$][A-Za-z0-9_$]*\b)"
    r"|(?P<punct>[{}()[\]<>:=;,.?+\-*/@#&|!]+)",
    re.MULTILINE,
)


def set_run_font(run, name: str, size: float, color: RGBColor, bold: bool = False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold


def shade_paragraph(paragraph, fill: str = "F6F8FA"):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_bottom_border(paragraph, color: str = "2E74B5", size: str = "10"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("PAGE ")
    set_run_font(run, "Calibri", 8.5, COLORS["muted"])
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run = paragraph.add_run()
    field_run._r.extend([begin, instr, separate, end])
    set_run_font(field_run, "Calibri", 8.5, COLORS["muted"])


def build_excerpt() -> tuple[str, list[str]]:
    stream = ""
    included: list[str] = []
    for source in SOURCE_FILES:
        relative = source.relative_to(ROOT).as_posix()
        marker = f"/* FILE: {relative} */\n"
        block = marker + source.read_text(encoding="utf-8").rstrip() + "\n\n"
        if len(stream) >= CHARACTER_LIMIT:
            break
        take = min(CHARACTER_LIMIT - len(stream), len(block))
        stream += block[:take]
        included.append(relative)
    assert len(stream) == CHARACTER_LIMIT
    return stream, included


def token_color(kind: str, value: str) -> tuple[RGBColor, bool]:
    if kind == "comment":
        return COLORS["comment"], False
    if kind == "string":
        return COLORS["green"], False
    if kind == "number":
        return COLORS["purple"], False
    if kind == "word" and value in KEYWORDS:
        return COLORS["blue"], True
    if kind == "word" and value and value[0].isupper():
        return COLORS["orange"], False
    if kind == "punct" and ("<" in value or ">" in value):
        return COLORS["red"], False
    return COLORS["ink"], False


def add_code_line(doc: Document, line: str):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.08)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(9.2)
    paragraph.paragraph_format.keep_together = True
    shade_paragraph(paragraph)

    cursor = 0
    for match in TOKEN_RE.finditer(line):
        if match.start() > cursor:
            run = paragraph.add_run(line[cursor:match.start()])
            set_run_font(run, "Consolas", 7.6, COLORS["ink"])
        value = match.group(0)
        kind = match.lastgroup or ""
        color, bold = token_color(kind, value)
        run = paragraph.add_run(value)
        set_run_font(run, "Consolas", 7.6, color, bold)
        cursor = match.end()
    if cursor < len(line):
        run = paragraph.add_run(line[cursor:])
        set_run_font(run, "Consolas", 7.6, COLORS["ink"])


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.36)
    section.footer_distance = Inches(0.36)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading = doc.styles["Heading 1"]
    heading.font.name = "Calibri"
    heading._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    heading._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    heading.font.size = Pt(16)
    heading.font.color.rgb = RGBColor(46, 116, 181)
    heading.paragraph_format.space_before = Pt(18)
    heading.paragraph_format.space_after = Pt(10)

    header = section.header.paragraphs[0]
    header.paragraph_format.space_after = Pt(0)
    header.paragraph_format.tab_stops.add_tab_stop(Inches(6.7))
    left = header.add_run("ZBX / SOURCE EXPORT")
    set_run_font(left, "Calibri", 8.5, COLORS["muted"], True)
    right = header.add_run("\tCURRENT WEBSITE / 2000 CHARACTERS")
    set_run_font(right, "Calibri", 8.5, COLORS["muted"])

    add_page_field(section.footer.paragraphs[0])


def main():
    excerpt, included = build_excerpt()
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(8)
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("CURRENT WEBSITE CODE")
    set_run_font(run, "Calibri", 22, COLORS["ink"], True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run("Source excerpt · exactly 2,000 characters")
    set_run_font(run, "Calibri", 11, COLORS["muted"])
    add_bottom_border(subtitle)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_before = Pt(8)
    meta.paragraph_format.space_after = Pt(10)
    meta.paragraph_format.line_spacing = 1.1
    lead = meta.add_run("Export order: ")
    set_run_font(lead, "Calibri", 9, COLORS["ink"], True)
    value = meta.add_run(" → ".join(included))
    set_run_font(value, "Calibri", 9, COLORS["muted"])

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(0)
    note.paragraph_format.space_after = Pt(8)
    run = note.add_run("Character count includes file markers and line breaks. Content after character 2,000 is intentionally omitted.")
    set_run_font(run, "Calibri", 8.5, COLORS["muted"])

    heading = doc.add_paragraph(style="Heading 1")
    heading.paragraph_format.space_before = Pt(6)
    heading.paragraph_format.space_after = Pt(6)
    heading.add_run("Source excerpt")

    for line in excerpt.split("\n"):
        add_code_line(doc, line)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "Current Website Code - 2000 Characters"
    doc.core_properties.subject = "ZBX portfolio source code excerpt"
    doc.core_properties.author = "Zhao Boxuan"
    doc.save(OUTPUT)
    print(f"OUTPUT={OUTPUT}")
    print(f"CODE_CHARACTERS={len(excerpt)}")
    print(f"FILES={included}")


if __name__ == "__main__":
    main()
