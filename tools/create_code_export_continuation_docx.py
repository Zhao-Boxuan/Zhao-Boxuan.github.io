from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
PREVIOUS = ROOT / "artifacts" / "ZBX_Website_Code_2000_Characters.docx"
OUTPUT = ROOT / "artifacts" / "ZBX_Website_Code_Continuation_3500_Characters.docx"
CHARACTER_LIMIT = 3500

SOURCE_FILES = [
    ROOT / "app" / "layout.tsx",
    ROOT / "app" / "page.tsx",
    ROOT / "styles" / "tokens.css",
    ROOT / "app" / "globals.css",
]

BLACK = RGBColor(0, 0, 0)


def set_font(run, name: str, size: float, bold: bool = False) -> None:
    run.font.name = name
    r_fonts = run._element.get_or_add_rPr().rFonts
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = BLACK
    run.bold = bold


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("PAGE ")
    set_font(run, "Arial", 8)

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run = paragraph.add_run()
    field_run._r.extend([begin, instr, separate, end])
    set_font(field_run, "Arial", 8)


def previous_excerpt() -> str:
    doc = Document(PREVIOUS)
    start = next(
        index for index, paragraph in enumerate(doc.paragraphs)
        if paragraph.text == "Source excerpt"
    ) + 1
    return "\n".join(paragraph.text for paragraph in doc.paragraphs[start:])


def find_overlap_anchor(old_excerpt: str, tokens_text: str) -> tuple[str, int]:
    lines = old_excerpt.splitlines()
    complete_lines = [line for line in lines if line.rstrip().endswith((";", "}", "*/"))]
    for line in reversed(complete_lines):
        position = tokens_text.find(line)
        if position >= 0:
            return line, position
    fallback = "  --color-sans-primary: #1a1a1a;"
    position = tokens_text.find(fallback)
    if position < 0:
        raise RuntimeError("Could not locate the previous excerpt cutoff in styles/tokens.css")
    return fallback, position


def build_excerpt() -> tuple[str, list[str], str]:
    old_excerpt = previous_excerpt()
    tokens_path = ROOT / "styles" / "tokens.css"
    globals_path = ROOT / "app" / "globals.css"
    tokens_text = tokens_path.read_text(encoding="utf-8").rstrip()
    globals_text = globals_path.read_text(encoding="utf-8").rstrip()
    anchor, anchor_position = find_overlap_anchor(old_excerpt, tokens_text)

    stream = (
        "/* FILE: styles/tokens.css - continued near previous cutoff */\n"
        + tokens_text[anchor_position:]
        + "\n\n/* FILE: app/globals.css */\n"
        + globals_text
        + "\n"
    )
    if len(stream) < CHARACTER_LIMIT:
        raise RuntimeError(f"Continuation stream has only {len(stream)} characters")
    return stream[:CHARACTER_LIMIT], ["styles/tokens.css", "app/globals.css"], anchor


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.34)
    section.footer_distance = Inches(0.34)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15

    title_style = doc.styles["Title"]
    title_style.font.name = "Arial"
    title_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title_style.font.size = Pt(21)
    title_style.font.bold = True
    title_style.font.color.rgb = BLACK
    title_style_p_pr = title_style._element.get_or_add_pPr()
    for border in title_style_p_pr.findall(qn("w:pBdr")):
        title_style_p_pr.remove(border)

    header = section.header.paragraphs[0]
    header.paragraph_format.space_after = Pt(0)
    run = header.add_run("ZBX  SOURCE EXPORT  CONTINUATION")
    set_font(run, "Arial", 8, True)

    add_page_field(section.footer.paragraphs[0])


def add_code_line(doc: Document, line: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.08)
    paragraph.paragraph_format.right_indent = Inches(0.04)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(10.2)
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run(line)
    set_font(run, "Consolas", 8.2)


def main() -> None:
    excerpt, included, anchor = build_excerpt()
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_before = Pt(6)
    title.paragraph_format.space_after = Pt(3)
    title_p_pr = title._p.get_or_add_pPr()
    for border in title_p_pr.findall(qn("w:pBdr")):
        title_p_pr.remove(border)
    title.add_run("Current Website Code Continuation")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(9)
    run = subtitle.add_run("Source continuation with exactly 3,500 characters")
    set_font(run, "Arial", 10.5)

    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(5)
    run = intro.add_run(
        "This document resumes near the final complete line of the previous "
        "2,000-character export and continues through the current website source."
    )
    set_font(run, "Arial", 9)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(11)
    run = meta.add_run("Export order  " + "  >  ".join(included))
    set_font(run, "Arial", 8.5)

    heading = doc.add_paragraph()
    heading.paragraph_format.space_before = Pt(0)
    heading.paragraph_format.space_after = Pt(5)
    run = heading.add_run("Source continuation")
    set_font(run, "Arial", 14, True)

    for line in excerpt.split("\n"):
        add_code_line(doc, line)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "Current Website Code Continuation 3500 Characters"
    doc.core_properties.subject = "ZBX portfolio source code continuation"
    doc.core_properties.author = "Zhao Boxuan"
    doc.save(OUTPUT)

    with ZipFile(OUTPUT) as archive:
        zip_error = archive.testzip()
    if zip_error is not None:
        raise RuntimeError(f"Invalid DOCX member: {zip_error}")

    print(f"OUTPUT={OUTPUT}")
    print(f"CODE_CHARACTERS={len(excerpt)}")
    print(f"FILES={included}")
    print(f"OVERLAP_ANCHOR={anchor}")


if __name__ == "__main__":
    main()
